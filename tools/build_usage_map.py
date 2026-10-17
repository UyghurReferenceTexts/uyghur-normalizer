#!/usr/bin/env python3
"""
build_usage_map.py — draft normalizer/usage_map.tsv from the توغرىسى entries
of the 6-volume izahliq lughet [2].

[2] lists common misuses as their own headwords, written

    X (توغرىسى: Y) ← Y          e.g.  ئاپارات (توغرىسى: ئاپپارات)

Each pair becomes a rule X -> Y unless one of these holds (first match wins):

  EXCLUDE  REAL_WORD     X is a headword of [4] or [1], or a word inside one
                         (گۈندە in ئارا گۈندە): a rule would corrupt a real
                         word (ئال "take", ئاپپاق "very white")
  EXCLUDE  SHORT         X has 3 letters or fewer: too likely to over-match
  EXCLUDE  NO_CHANGE     X and Y are the same after normalization
  EXCLUDE  MAP_CONFLICT  X is the output of a rule in another map, so this
                         rule would undo that one
  HOLD     MULTI_TARGET  Y names more than one form, or X maps to two Ys
  HOLD     UNATTESTED    Y, after the era and loan passes, is not in [4] or
                         [1]. [2] predates the 2009 reform, so its spelling
                         of Y cannot be trusted on its own

A pair listed in tools/usage_target_overrides.tsv takes its correct form from
there: for when [2] gives a correct but pre-2009 spelling that the passes do
not fix (چىڭگىلىماچ: [2] جىڭگىلماچ, [1] جىڭگىلىماچ).

The usage pass runs before the other passes, so a pre-2009 spelling in Y is
corrected by the era and loan passes automatically.

A rule gets `nosfx` when X is the beginning of a real word or of Y: without
it, the suffix engine would read the rest of that word as a suffix.

    python tools/build_usage_map.py --vol1990 izahliq_lughet_1990s_v1.xlsx ... \
        [--reviewed toghrisi_pairs_1990s.xlsx] [--xlsx usage_map_review.xlsx]

Requires openpyxl (tooling only).
"""
import argparse, collections, csv, gzip, os, re, sys, unicodedata
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from uyghur_normalizer import load_map, build_pattern, process_text, Report  # noqa: E402

PAIR = re.compile(r'^\s*(.+?)\s*\(\s*توغرىسى\s*[:：]\s*(.+?)\s*\)')
LETTERS = re.compile(r'[\u0620-\u06ff]')
SNAP = os.path.join(ROOT, 'tests', 'data', 'lexicon_snapshot.tsv.gz')
OUT = os.path.join(ROOT, 'normalizer', 'usage_map.tsv')
OVERRIDES = os.path.join(ROOT, 'tools', 'usage_target_overrides.tsv')


def clean(s):
    s = unicodedata.normalize('NFC', s)
    s = re.sub(r'[\d①-⑳]+', '', s)
    return re.sub(r'\s+', ' ', s).strip(' .،')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--vol1990', nargs=6, required=True, help='the six volumes, in order')
    ap.add_argument('--reviewed', help='reviewed pairs workbook; its Note column is carried over')
    ap.add_argument('--xlsx', default='usage_map_review.xlsx')
    ap.add_argument('--out', default=OUT)
    a = ap.parse_args()

    with gzip.open(SNAP, 'rt', encoding='utf-8') as fh:
        lex = {}
        for r in csv.DictReader((l for l in fh if not l.startswith('#')), delimiter='\t',
                                quoting=csv.QUOTE_NONE):
            if r['form'] not in lex or r['status'] == 'verified':
                lex[r['form']] = r
    forms = sorted(lex)
    parts = {}                     # every word inside a multi-word or hyphenated headword
    for f, r in lex.items():
        for t in re.split(r'[\s\-]+', f):
            if t and t != f and (t not in parts or r['status'] == 'verified'):
                parts[t] = r

    nd = os.path.join(ROOT, 'normalizer')
    other = {m: load_map(os.path.join(nd, m), m, quiet=True)
             for m in ('calque_map.tsv', 'era_map.tsv', 'imla_loan_map.tsv')}
    outputs = {rhs for rules in other.values() for _l, rhs, *_ in rules}
    ortho = [build_pattern(other[m]) for m in ('era_map.tsv', 'imla_loan_map.tsv')]

    notes = {}
    if a.reviewed:
        ws = openpyxl.load_workbook(a.reviewed, read_only=True).worksheets[0]
        head = col = None
        for r in ws.iter_rows(values_only=True):
            if head is None:                         # the header may not be the first row
                cells = [str(x) for x in r]
                for name in ('Note', 'your earlier note'):
                    if name in cells:
                        head, col = cells, cells.index(name)
                continue
            if r[0] and r[col]:
                notes[r[0]] = str(r[col]).split(' Target overridden')[0]

    pairs = []
    for v, path in enumerate(a.vol1990, 1):
        ws = openpyxl.load_workbook(path, read_only=True)[f'Dictionary Vol{v}']
        for r in ws.iter_rows(min_row=2, values_only=True):
            m = PAIR.match(str(r[1] or ''))
            if m:
                pairs.append(dict(row_id=r[0], wrong=clean(m.group(1)), right1990=clean(m.group(2))))

    overrides = {}
    if os.path.exists(OVERRIDES):
        with open(OVERRIDES, encoding='utf-8') as fh:
            for r in csv.DictReader((l for l in fh if not l.startswith('#')), delimiter='\t'):
                overrides[r['variant']] = r['target']
    for p in pairs:
        if p['wrong'] in overrides:
            p['printed1990'], p['right1990'] = p['right1990'], overrides[p['wrong']]

    targets = collections.defaultdict(set)
    for p in pairs:
        targets[p['wrong']].add(p['right1990'])

    def prefix_of_word(w):
        import bisect
        i = bisect.bisect_right(forms, w)
        return i < len(forms) and forms[i].startswith(w) and forms[i] != w

    rules = []
    for p in pairs:
        w, y = p['wrong'], p['right1990']
        y_norm = process_text(y, ortho, Report())
        p['right_2009'] = y_norm
        ev = lex.get(y_norm)
        p['evidence'] = f"{ev['source']} {ev['row_id']}" if ev else ''
        p['evidence_scan'] = ev['scan'] if ev else ''
        p['evidence_verified'] = ('yes' if ev['status'] == 'verified' else 'no') if ev else ''
        p['note'] = notes.get(p['row_id'], '')
        if 'printed1990' in p:
            p['note'] = (p['note'] + ' ' if p['note'] else '') + \
                f"Target overridden (tools/usage_target_overrides.tsv); [2] prints {p['printed1990']}."

        p['flags'] = ''
        if w in lex:
            d = ('EXCLUDE', 'REAL_WORD', f"{lex[w]['source']} {lex[w]['row_id']}")
        elif w in parts:
            d = ('EXCLUDE', 'REAL_WORD', f"inside a headword: {parts[w]['source']} {parts[w]['row_id']}")
        elif len(LETTERS.findall(w)) <= 3:
            d = ('EXCLUDE', 'SHORT', '')
        elif process_text(w, ortho, Report()) == y_norm:
            d = ('EXCLUDE', 'NO_CHANGE', '')
        elif w in outputs:
            d = ('EXCLUDE', 'MAP_CONFLICT', 'is the output of an existing rule')
        elif re.search('[،,;؛/]', y) or len(targets[w]) > 1:
            d = ('HOLD', 'MULTI_TARGET', ' | '.join(sorted(targets[w])))
        elif not ev:
            d = ('HOLD', 'UNATTESTED', '')
        else:
            d = ('INCLUDE', '', '')
            if y_norm.startswith(w) or prefix_of_word(w):
                p['flags'] = 'nosfx'
        p['decision'], p['reason'], p['detail'] = d
        if d[0] == 'INCLUDE':
            rules.append(p)

    seen, uniq = set(), []
    for p in rules:
        if p['wrong'] not in seen:
            seen.add(p['wrong'])
            uniq.append(p)
    uniq.sort(key=lambda p: (-len(p['wrong']), p['wrong']))
    with open(a.out, 'w', encoding='utf-8', newline='\n') as fh:
        fh.write('# usage_map.tsv — common misuses and their correct forms (OPT-IN: --usage)\n'
                 '# PASS 0, before calque -> era -> imla_loan, so pre-2009 spellings in the\n'
                 '# replacements are normalized by passes 2-3.\n'
                 '# Source: the توغرىسى entries of [2], each target attested in [4] or [1]\n'
                 '# after normalization. Row ids per rule: normalizer/provenance.tsv.\n'
                 '# Generated by tools/build_usage_map.py; review there, not here.\n'
                 f'# {len(uniq)} rules. variant<TAB>standard[<TAB>flags]\n')
        for p in uniq:
            fh.write(f"{p['wrong']}\t{p['right1990']}" + (f"\t{p['flags']}" if p['flags'] else '') + '\n')

    with open(os.path.join(os.path.dirname(a.out), 'usage_map_origin.tsv'), 'w',
              encoding='utf-8', newline='\n') as fh:
        fh.write('# usage_map_origin.tsv — the [2] row each usage rule comes from; generated\n')
        fh.write('variant\trow_id_1990s\n')
        for p in uniq:
            fh.write(f"{p['wrong']}\t{p['row_id']}\n")

    write_xlsx(a.xlsx, pairs)
    c = collections.Counter((p['decision'], p['reason']) for p in pairs)
    for k in sorted(c):
        print(f'{k[0]:8} {k[1]:13} {c[k]}')
    print(f"{len(uniq)} rules ({sum(1 for p in uniq if p['flags'])} nosfx) -> {a.out}")


def write_xlsx(path, pairs):
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'usage pairs'
    cols = [('row_id_1990s', 'row_id', 14), ('wrong form', 'wrong', 22), ('correct form [2]', 'right1990', 22),
            ('after era/loan passes', 'right_2009', 22), ('decision', 'decision', 10), ('reason', 'reason', 15),
            ('detail', 'detail', 26), ('flags', 'flags', 8), ('target attested in', 'evidence', 16),
            ('scan', 'evidence_scan', 15), ('verified', 'evidence_verified', 9), ('your earlier note', 'note', 34),
            ('your verdict', None, 30)]
    f = Font(name='Arial', size=11)
    ws.append([c[0] for c in cols])
    for p in pairs:
        ws.append([p[k] if k else '' for _h, k, _w in cols])
    fills = {'INCLUDE': 'E2EFDA', 'HOLD': 'FFF2CC', 'EXCLUDE': 'F2F2F2'}
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = f
        row[4].fill = PatternFill('solid', fgColor=fills[row[4].value])
    for i, (_h, _k, w) in enumerate(cols, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        ws.cell(1, i).font = Font(name='Arial', size=11, bold=True)
    ws.cell(1, len(cols)).fill = PatternFill('solid', fgColor='DDEBF7')
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    lg = wb.create_sheet('legend')
    for line in [
        ('INCLUDE', 'Goes into usage_map.tsv.'),
        ('HOLD', 'Kept out until you decide. Write keep / drop in "your verdict".'),
        ('EXCLUDE', 'Kept out automatically. Overrule only with a reason.'),
        ('REAL_WORD', 'The wrong form is a headword of [4] or [1], or a word inside one; a rule would corrupt it.'),
        ('SHORT', 'Three letters or fewer; too likely to match inside other words.'),
        ('NO_CHANGE', 'After normalization the two forms are identical.'),
        ('MAP_CONFLICT', 'The wrong form is what another map produces; the rule would undo it.'),
        ('MULTI_TARGET', 'More than one correct form is given; pick one.'),
        ('UNATTESTED', 'The correct form is not in [4] or [1], even after normalization.'),
        ('nosfx', 'The wrong form is the start of a real word, so suffixed forms are left alone.'),
        ('your verdict', 'The only column to edit (blue header).'),
    ]:
        lg.append(line)
    for row in lg.iter_rows():
        for cell in row:
            cell.font = Font(name='Arial', size=11, bold=cell.column == 1)
    lg.column_dimensions['A'].width = 16
    lg.column_dimensions['B'].width = 90
    for ws_ in (ws,):
        ws_.sheet_view.rightToLeft = False
    wb.save(path)


if __name__ == '__main__':
    main()
